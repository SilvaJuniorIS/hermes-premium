from __future__ import annotations

import json
import os
import secrets
import threading
import time
from email.utils import parseaddr
from io import BytesIO
from base64 import urlsafe_b64decode, urlsafe_b64encode
from datetime import datetime
from hashlib import pbkdf2_hmac
from hmac import compare_digest
from pathlib import Path
from typing import Any

from fastapi import Cookie, Depends, FastAPI, Form, HTTPException
from fastapi.responses import HTMLResponse, RedirectResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles
from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from openpyxl.styles import Alignment, Font, PatternFill
from pydantic import BaseModel, Field

from src.hermes_db import (
    ensure_admin_user,
    get_user_by_id,
    get_user_by_username,
    init_db,
    load_licitacao_detail,
    load_recent_api_runs,
    load_recent_licitacoes,
    mark_user_login,
    update_licitacao_comercial,
)
from src.main import COLETA_GERAL, DEFAULT_CONFIG, run_pipeline, run_pipeline_coleta_geral
from src.notifier import enviar_planilha_oportunidades


app = FastAPI(title="HERMES - Inteligencia em Licitacoes Publicas")
PERFIS_PATH = Path("config/perfis_negocio.json")
SCHEDULE_PATH = Path("config/agendamento.json")
SESSION_COOKIE = "hermes_session"
SESSION_TOKENS: dict[str, int] = {}
SCHEDULER_LOCK = threading.Lock()
SCHEDULER_STARTED = False

Path("docs/assets").mkdir(parents=True, exist_ok=True)
app.mount("/assets", StaticFiles(directory="docs/assets"), name="assets")


class RunRequest(BaseModel):
    perfil: str = "limpeza_higiene"
    config: dict[str, Any] | None = None


class PerfilRequest(BaseModel):
    id: str
    nome_exibicao: str
    descricao: str = ""
    keywords: list[str] = Field(default_factory=list)
    keywords_fortes: list[str] = Field(default_factory=list)
    termos_positivos: list[str] = Field(default_factory=list)
    termos_negativos: list[str] = Field(default_factory=list)
    valor_minimo_interesse: float = 0
    valor_atrativo: float = 50000
    valor_muito_atrativo: float = 200000
    score_relevancia_media: float = 50
    score_relevancia_alta: float = 75


class ScheduleRequest(BaseModel):
    enabled: bool = False
    perfil: str = "limpeza_higiene"
    time: str = "08:00"
    config: dict[str, Any] | None = None


class EmailExportRequest(BaseModel):
    destinatario: str
    limit: int = 500
    perfil: str | None = None
    classificacao: str | None = None
    estado: str | None = None
    q: str | None = None
    valor_min: float | None = None
    valor_max: float | None = None
    status_comercial: str | None = None
    favorito: bool | None = None
    prazo: str | None = "futuras"


class ComercialUpdateRequest(BaseModel):
    status_comercial: str = "novo"
    favorito: bool = False
    anotacoes: str = ""


def _hash_password(password: str, salt: bytes | None = None) -> str:
    salt = salt or os.urandom(16)
    digest = pbkdf2_hmac("sha256", password.encode("utf-8"), salt, 120000)
    return (
        urlsafe_b64encode(salt).decode("ascii")
        + "$"
        + urlsafe_b64encode(digest).decode("ascii")
    )


def _verify_password(password: str, stored_hash: str) -> bool:
    try:
        salt_text, digest_text = stored_hash.split("$", 1)
        salt = urlsafe_b64decode(salt_text.encode("ascii"))
        expected = urlsafe_b64decode(digest_text.encode("ascii"))
    except (ValueError, TypeError):
        return False

    actual = pbkdf2_hmac("sha256", password.encode("utf-8"), salt, 120000)
    return compare_digest(actual, expected)


def _login_page(error: str = "") -> str:
    error_html = f'<div class="error">{error}</div>' if error else ""
    return f"""
<!DOCTYPE html>
<html lang="pt-BR">
<head>
    <meta charset="utf-8">
    <meta name="viewport" content="width=device-width, initial-scale=1">
    <title>Login - HERMES</title>
    <style>
        * {{ box-sizing: border-box; }}
        body {{
            margin: 0;
            min-height: 100vh;
            display: grid;
            place-items: center;
            background: #0a2342;
            color: #f5f7fa;
            font-family: Inter, Roboto, Arial, Helvetica, sans-serif;
        }}
        form {{
            width: min(420px, calc(100vw - 32px));
            border: 1px solid rgba(245, 247, 250, 0.14);
            border-radius: 8px;
            background: #102f55;
            padding: 24px;
            box-shadow: 0 18px 48px rgba(0, 0, 0, 0.24);
        }}
        h1 {{
            margin: 0 0 8px;
            font-size: 24px;
            font-family: Poppins, Montserrat, Inter, Arial, sans-serif;
        }}
        h1 span {{
            color: #f7931e;
        }}
        .tagline {{
            margin: 0 0 18px;
            color: #cbd5e1;
            font-size: 13px;
        }}
        label {{
            display: grid;
            gap: 6px;
            margin-bottom: 14px;
            color: #9fb0bf;
            font-size: 13px;
        }}
        input, button {{
            height: 40px;
            border: 1px solid rgba(245, 247, 250, 0.16);
            border-radius: 6px;
            background: #0a2342;
            color: #f5f7fa;
            padding: 0 12px;
            font-size: 15px;
        }}
        button {{
            width: 100%;
            border-color: #d97706;
            background: #f7931e;
            cursor: pointer;
            font-weight: 700;
        }}
        .error {{
            margin-bottom: 14px;
            border: 1px solid #8f3939;
            border-radius: 6px;
            background: rgba(249, 115, 115, 0.12);
            color: #f97373;
            padding: 10px;
            font-size: 14px;
        }}
        p {{
            margin: 14px 0 0;
            color: #9fb0bf;
            font-size: 13px;
        }}
    </style>
</head>
<body>
    <form method="post" action="/login">
        <h1>HER<span>MES</span></h1>
        <p class="tagline">Inteligencia em Licitacoes Publicas</p>
        {error_html}
        <label>
            Usuario
            <input name="username" autocomplete="username" required autofocus>
        </label>
        <label>
            Senha
            <input name="password" type="password" autocomplete="current-password" required>
        </label>
        <button type="submit">Entrar</button>
        <p>Usuario inicial: admin. Senha inicial: admin123.</p>
    </form>
</body>
</html>
"""


def _current_user(hermes_session: str | None = Cookie(default=None)) -> dict[str, Any]:
    if not hermes_session:
        raise HTTPException(status_code=401, detail="Nao autenticado")

    user_id = SESSION_TOKENS.get(hermes_session)
    if not user_id:
        raise HTTPException(status_code=401, detail="Sessao invalida")

    user = get_user_by_id(user_id)
    if not user:
        raise HTTPException(status_code=401, detail="Usuario inativo ou inexistente")

    return user


def _load_perfis() -> dict[str, Any]:
    if not PERFIS_PATH.exists():
        return {}
    try:
        return json.loads(PERFIS_PATH.read_text(encoding="utf-8-sig"))
    except json.JSONDecodeError as exc:
        raise HTTPException(status_code=500, detail=f"Erro ao ler perfis: {exc}") from exc


def _save_perfis(perfis: dict[str, Any]) -> None:
    PERFIS_PATH.parent.mkdir(parents=True, exist_ok=True)
    PERFIS_PATH.write_text(
        json.dumps(perfis, ensure_ascii=False, indent=4),
        encoding="utf-8",
    )


def _load_schedule() -> dict[str, Any]:
    default = {
        "enabled": False,
        "perfil": "limpeza_higiene",
        "time": "08:00",
        "config": {},
        "last_run_date": "",
        "last_run_status": "",
    }
    if not SCHEDULE_PATH.exists():
        return default
    try:
        data = json.loads(SCHEDULE_PATH.read_text(encoding="utf-8-sig"))
    except (OSError, json.JSONDecodeError):
        return default

    return {**default, **data}


def _save_schedule(data: dict[str, Any]) -> None:
    SCHEDULE_PATH.parent.mkdir(parents=True, exist_ok=True)
    SCHEDULE_PATH.write_text(
        json.dumps(data, ensure_ascii=False, indent=4),
        encoding="utf-8",
    )


def _valid_schedule_time(value: str) -> str:
    try:
        datetime.strptime(value, "%H:%M")
    except ValueError as exc:
        raise HTTPException(status_code=400, detail="Horario deve usar HH:MM") from exc
    return value


def _scheduler_loop() -> None:
    while True:
        try:
            schedule = _load_schedule()
            now = datetime.now()
            today = now.date().isoformat()
            if (
                schedule.get("enabled")
                and schedule.get("time") == now.strftime("%H:%M")
                and schedule.get("last_run_date") != today
            ):
                with SCHEDULER_LOCK:
                    schedule = _load_schedule()
                    if schedule.get("last_run_date") == today:
                        time.sleep(30)
                        continue

                    perfil = schedule.get("perfil") or "limpeza_higiene"
                    config = {**DEFAULT_CONFIG, **(schedule.get("config") or {})}
                    try:
                        if perfil == COLETA_GERAL:
                            result = run_pipeline_coleta_geral(config)
                            stats = result["stats"]
                            ok = stats.get("perfis_ok", 0)
                            total = stats.get("perfis_total", 0)
                            schedule["last_run_status"] = (
                                f"finished: {stats.get('salvas', 0)} registros "
                                f"(coleta geral {ok}/{total} perfis)"
                            )
                        else:
                            result = run_pipeline(perfil, config)
                            schedule["last_run_status"] = (
                                f"finished: {result['stats'].get('salvas', 0)} registros"
                            )
                    except Exception as exc:  # pragma: no cover - defensive background loop
                        schedule["last_run_status"] = f"error: {exc}"
                    schedule["last_run_date"] = today
                    _save_schedule(schedule)
            time.sleep(30)
        except Exception:
            time.sleep(60)


def _start_scheduler_once() -> None:
    global SCHEDULER_STARTED
    if SCHEDULER_STARTED:
        return
    thread = threading.Thread(target=_scheduler_loop, daemon=True)
    thread.start()
    SCHEDULER_STARTED = True


def _perfil_payload(req: PerfilRequest) -> dict[str, Any]:
    return {
        "nome_exibicao": req.nome_exibicao,
        "descricao": req.descricao,
        "keywords": req.keywords,
        "keywords_fortes": req.keywords_fortes,
        "termos_positivos": req.termos_positivos,
        "termos_negativos": req.termos_negativos,
        "valor_minimo_interesse": req.valor_minimo_interesse,
        "valor_atrativo": req.valor_atrativo,
        "valor_muito_atrativo": req.valor_muito_atrativo,
        "score_relevancia_media": req.score_relevancia_media,
        "score_relevancia_alta": req.score_relevancia_alta,
    }


@app.on_event("startup")
def startup() -> None:
    init_db()
    ensure_admin_user(_hash_password)
    _start_scheduler_once()


@app.get("/", response_class=HTMLResponse)
def dashboard(hermes_session: str | None = Cookie(default=None)):
    user_id = SESSION_TOKENS.get(hermes_session or "")
    if not user_id or not get_user_by_id(user_id):
        return RedirectResponse("/login", status_code=303)

    dashboard_path = Path("dashboard.html")
    return HTMLResponse(dashboard_path.read_text(encoding="utf-8"))


@app.get("/login", response_class=HTMLResponse)
def login_form() -> str:
    init_db()
    ensure_admin_user(_hash_password)
    return _login_page()


@app.post("/login")
def login(username: str = Form(...), password: str = Form(...)) -> RedirectResponse:
    init_db()
    ensure_admin_user(_hash_password)
    user = get_user_by_username(username.strip())
    if not user or not _verify_password(password, user["password_hash"]):
        return HTMLResponse(_login_page("Usuario ou senha invalidos"), status_code=401)

    token = secrets.token_urlsafe(32)
    SESSION_TOKENS[token] = int(user["id"])
    mark_user_login(int(user["id"]))

    response = RedirectResponse("/", status_code=303)
    response.set_cookie(
        SESSION_COOKIE,
        token,
        httponly=True,
        samesite="lax",
        max_age=60 * 60 * 8,
    )
    return response


@app.post("/logout")
def logout(hermes_session: str | None = Cookie(default=None)) -> RedirectResponse:
    if hermes_session:
        SESSION_TOKENS.pop(hermes_session, None)

    response = RedirectResponse("/login", status_code=303)
    response.delete_cookie(SESSION_COOKIE)
    return response


@app.post("/run")
def run(req: RunRequest, _: dict[str, Any] = Depends(_current_user)) -> dict[str, Any]:
    config = {**DEFAULT_CONFIG, **(req.config or {})}
    if req.perfil == COLETA_GERAL:
        result = run_pipeline_coleta_geral(config)
    else:
        result = run_pipeline(req.perfil, config)
    return {
        "status": result["status"],
        "run_id": result["run_id"],
        "stats": result["stats"],
    }


@app.get("/licitacoes")
def listar(
    limit: int = 100,
    perfil: str | None = None,
    classificacao: str | None = None,
    estado: str | None = None,
    q: str | None = None,
    valor_min: float | None = None,
    valor_max: float | None = None,
    status_comercial: str | None = None,
    favorito: bool | None = None,
    prazo: str | None = "futuras",
    _: dict[str, Any] = Depends(_current_user),
) -> list[dict[str, Any]]:
    return load_recent_licitacoes(
        limit=limit,
        perfil=perfil,
        classificacao=classificacao,
        estado=estado,
        q=q,
        valor_min=valor_min,
        valor_max=valor_max,
        status_comercial=status_comercial,
        favorito=favorito,
        prazo=_normalize_prazo(prazo),
    )


def _build_licitacoes_workbook(rows: list[dict[str, Any]]) -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = "Oportunidades"
    ws.append(["HERMES - Inteligencia em Licitacoes Publicas"])
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=21)
    title_cell = ws.cell(1, 1)
    title_cell.font = Font(bold=True, color="F7931E", size=15)
    title_cell.fill = PatternFill("solid", fgColor="0A2342")
    title_cell.alignment = Alignment(horizontal="center")
    ws.append(["Dados publicos transformados em oportunidades estrategicas"])
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=21)
    subtitle_cell = ws.cell(2, 1)
    subtitle_cell.font = Font(color="2D3748", italic=True)
    subtitle_cell.alignment = Alignment(horizontal="center")
    headers = [
        "acao_recomendada",
        "status_comercial",
        "favorito",
        "perfil",
        "score",
        "classificacao",
        "objeto",
        "estado",
        "municipio",
        "orgao",
        "valor_estimado",
        "oportunidade",
        "data_abertura",
        "janela_comercial",
        "keyword",
        "motivos_score",
        "anotacoes",
        "primeiro_registro",
        "ultima_leitura",
        "link",
        "pncp_id",
    ]
    ws.append(headers)
    header_row = 3
    header_fill = PatternFill("solid", fgColor="0A2342")
    header_font = Font(bold=True, color="FFFFFF")
    for cell in ws[header_row]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for item in rows:
        ws.append(
            [
                _acao_recomendada(item),
                item.get("status_comercial") or "novo",
                "sim" if item.get("favorito") else "nao",
                item.get("perfil"),
                item.get("score"),
                item.get("classificacao"),
                item.get("objeto"),
                item.get("estado"),
                item.get("municipio"),
                item.get("orgao"),
                item.get("valor_estimado_num"),
                item.get("oportunidade"),
                item.get("data_abertura"),
                _janela_comercial(item.get("data_abertura")),
                item.get("keyword"),
                "; ".join(item.get("score_motivos") or []),
                item.get("anotacoes"),
                item.get("first_seen_at"),
                item.get("last_seen_at"),
                item.get("link"),
                item.get("pncp_id"),
            ]
        )
    for column_index, column in enumerate(ws.columns, start=1):
        letter = get_column_letter(column_index)
        ws.column_dimensions[letter].width = min(
            max(len(str(cell.value or "")) for cell in column) + 2,
            55,
        )
    ws.freeze_panes = "A4"
    return wb


def _acao_recomendada(item: dict[str, Any]) -> str:
    score = float(item.get("score") or 0)
    valor = float(item.get("valor_estimado_num") or 0)
    if score >= 85 and valor >= 200000:
        return "Priorizar contato e analise do edital"
    if score >= 75:
        return "Avaliar proposta nesta semana"
    if score >= 50:
        return "Monitorar e validar aderencia"
    return "Manter no radar"


def _janela_comercial(value: Any) -> str:
    if not value:
        return "Prazo nao informado"
    try:
        normalized = str(value).replace("Z", "").replace("T", " ")[:19]
        abertura = datetime.fromisoformat(normalized)
    except ValueError:
        return "Data de abertura invalida"

    dias = (abertura.date() - datetime.now().date()).days
    if dias < 0:
        return "Abertura ja passou"
    if dias == 0:
        return "Abre hoje"
    if dias <= 7:
        return f"{dias} dias: decisao imediata"
    if dias <= 21:
        return f"{dias} dias: janela boa para proposta"
    return f"{dias} dias: monitorar e preparar abordagem"


def _valid_email(value: str) -> str:
    _, email = parseaddr(value.strip())
    if not email or "@" not in email or "." not in email.rsplit("@", 1)[-1]:
        raise HTTPException(status_code=400, detail="E-mail do destinatario invalido")
    return email


def _describe_export_filters(req: EmailExportRequest) -> str:
    filtros = []
    if req.perfil:
        filtros.append(f"perfil={req.perfil}")
    if req.classificacao:
        filtros.append(f"classificacao={req.classificacao}")
    if req.estado:
        filtros.append(f"estado={req.estado}")
    if req.q:
        filtros.append(f"texto={req.q}")
    if req.valor_min is not None:
        filtros.append(f"valor_min={req.valor_min}")
    if req.valor_max is not None:
        filtros.append(f"valor_max={req.valor_max}")
    if req.status_comercial:
        filtros.append(f"status_comercial={req.status_comercial}")
    if req.favorito is not None:
        filtros.append(f"favorito={req.favorito}")
    if req.prazo:
        filtros.append(f"prazo={req.prazo}")
    filtros.append(f"limite={req.limit}")
    return ", ".join(filtros)


def _normalize_commercial_status(value: str) -> str:
    allowed = {"novo", "em_analise", "proposta", "descartado", "monitorar"}
    normalized = value.strip().lower()
    if normalized not in allowed:
        raise HTTPException(status_code=400, detail="Status comercial invalido")
    return normalized


def _normalize_prazo(value: str | None) -> str | None:
    if not value or value == "todas":
        return None
    if value not in {"futuras", "vencidas"}:
        raise HTTPException(status_code=400, detail="Filtro de prazo invalido")
    return value


@app.get("/licitacoes/export.xlsx")
def exportar_licitacoes_excel(
    limit: int = 500,
    perfil: str | None = None,
    classificacao: str | None = None,
    estado: str | None = None,
    q: str | None = None,
    valor_min: float | None = None,
    valor_max: float | None = None,
    status_comercial: str | None = None,
    favorito: bool | None = None,
    prazo: str | None = "futuras",
    _: dict[str, Any] = Depends(_current_user),
) -> StreamingResponse:
    rows = load_recent_licitacoes(
        limit=limit,
        perfil=perfil,
        classificacao=classificacao,
        estado=estado,
        q=q,
        valor_min=valor_min,
        valor_max=valor_max,
        status_comercial=status_comercial,
        favorito=favorito,
        prazo=_normalize_prazo(prazo),
    )
    wb = _build_licitacoes_workbook(rows)
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    filename = f"hermes-oportunidades-{datetime.now().date().isoformat()}.xlsx"
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@app.get("/licitacoes/export/local")
def exportar_licitacoes_excel_local(
    limit: int = 500,
    perfil: str | None = None,
    classificacao: str | None = None,
    estado: str | None = None,
    q: str | None = None,
    valor_min: float | None = None,
    valor_max: float | None = None,
    status_comercial: str | None = None,
    favorito: bool | None = None,
    prazo: str | None = "futuras",
    _: dict[str, Any] = Depends(_current_user),
) -> dict[str, Any]:
    rows = load_recent_licitacoes(
        limit=limit,
        perfil=perfil,
        classificacao=classificacao,
        estado=estado,
        q=q,
        valor_min=valor_min,
        valor_max=valor_max,
        status_comercial=status_comercial,
        favorito=favorito,
        prazo=_normalize_prazo(prazo),
    )
    wb = _build_licitacoes_workbook(rows)
    output_dir = Path("output")
    output_dir.mkdir(parents=True, exist_ok=True)
    filename = f"hermes-oportunidades-{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    output_path = output_dir / filename
    wb.save(output_path)
    return {
        "status": "saved",
        "total": len(rows),
        "filename": filename,
        "path": str(output_path.resolve()),
    }


@app.post("/licitacoes/export/email")
def enviar_licitacoes_excel_email(
    req: EmailExportRequest,
    _: dict[str, Any] = Depends(_current_user),
) -> dict[str, Any]:
    destinatario = _valid_email(req.destinatario)
    rows = load_recent_licitacoes(
        limit=min(max(req.limit, 1), 1000),
        perfil=req.perfil,
        classificacao=req.classificacao,
        estado=req.estado,
        q=req.q,
        valor_min=req.valor_min,
        valor_max=req.valor_max,
        status_comercial=req.status_comercial,
        favorito=req.favorito,
        prazo=_normalize_prazo(req.prazo),
    )
    if not rows:
        raise HTTPException(status_code=400, detail="Nao ha oportunidades para enviar")

    wb = _build_licitacoes_workbook(rows)
    output = BytesIO()
    wb.save(output)
    arquivo_bytes = output.getvalue()
    filename = f"hermes-oportunidades-{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"

    try:
        enviar_planilha_oportunidades(
            destinatario=destinatario,
            arquivo_bytes=arquivo_bytes,
            nome_arquivo=filename,
            total=len(rows),
            filtros=_describe_export_filters(req),
            destaques=rows,
        )
    except Exception as exc:
        raise HTTPException(status_code=502, detail=f"Erro ao enviar e-mail: {exc}") from exc

    return {
        "status": "sent",
        "total": len(rows),
        "filename": filename,
        "destinatario": destinatario,
    }


@app.get("/licitacao")
def detalhe_licitacao(
    pncp_id: str,
    _: dict[str, Any] = Depends(_current_user),
) -> dict[str, Any]:
    detail = load_licitacao_detail(pncp_id)
    if not detail:
        raise HTTPException(status_code=404, detail="Licitacao nao encontrada")
    return detail


@app.put("/licitacao/comercial/{pncp_id}")
def atualizar_licitacao_comercial(
    pncp_id: str,
    req: ComercialUpdateRequest,
    _: dict[str, Any] = Depends(_current_user),
) -> dict[str, Any]:
    detail = update_licitacao_comercial(
        pncp_id=pncp_id,
        status_comercial=_normalize_commercial_status(req.status_comercial),
        favorito=req.favorito,
        anotacoes=req.anotacoes.strip(),
    )
    if not detail:
        raise HTTPException(status_code=404, detail="Licitacao nao encontrada")
    return detail


@app.get("/runs")
def runs(_: dict[str, Any] = Depends(_current_user)) -> list[dict[str, Any]]:
    return load_recent_api_runs()


@app.get("/schedule")
def get_schedule(_: dict[str, Any] = Depends(_current_user)) -> dict[str, Any]:
    return _load_schedule()


@app.put("/schedule")
def update_schedule(
    req: ScheduleRequest,
    _: dict[str, Any] = Depends(_current_user),
) -> dict[str, Any]:
    data = _load_schedule()
    data.update(
        {
            "enabled": req.enabled,
            "perfil": req.perfil,
            "time": _valid_schedule_time(req.time),
            "config": req.config or {},
        }
    )
    _save_schedule(data)
    return data


@app.get("/health")
def health() -> dict[str, str]:
    init_db()
    return {"status": "ok"}


@app.get("/perfis")
def listar_perfis(_: dict[str, Any] = Depends(_current_user)) -> dict[str, Any]:
    return _load_perfis()


@app.post("/perfis")
def criar_perfil(
    req: PerfilRequest,
    _: dict[str, Any] = Depends(_current_user),
) -> dict[str, Any]:
    perfil_id = req.id.strip().lower().replace(" ", "_")
    if not perfil_id:
        raise HTTPException(status_code=400, detail="ID do perfil e obrigatorio")

    perfis = _load_perfis()
    if perfil_id in perfis:
        raise HTTPException(status_code=409, detail="Perfil ja existe")

    perfis[perfil_id] = _perfil_payload(req)
    _save_perfis(perfis)
    return {"id": perfil_id, "perfil": perfis[perfil_id]}


@app.put("/perfis/{perfil_id}")
def atualizar_perfil(
    perfil_id: str,
    req: PerfilRequest,
    _: dict[str, Any] = Depends(_current_user),
) -> dict[str, Any]:
    perfis = _load_perfis()
    perfil_id = perfil_id.strip().lower()
    if perfil_id not in perfis:
        raise HTTPException(status_code=404, detail="Perfil nao encontrado")

    novo_id = req.id.strip().lower().replace(" ", "_")
    if not novo_id:
        raise HTTPException(status_code=400, detail="ID do perfil e obrigatorio")
    if novo_id != perfil_id and novo_id in perfis:
        raise HTTPException(status_code=409, detail="Novo ID ja existe")

    payload = _perfil_payload(req)
    if novo_id != perfil_id:
        del perfis[perfil_id]
    perfis[novo_id] = payload
    _save_perfis(perfis)
    return {"id": novo_id, "perfil": payload}


@app.delete("/perfis/{perfil_id}")
def excluir_perfil(
    perfil_id: str,
    _: dict[str, Any] = Depends(_current_user),
) -> dict[str, str]:
    perfis = _load_perfis()
    perfil_id = perfil_id.strip().lower()
    if perfil_id not in perfis:
        raise HTTPException(status_code=404, detail="Perfil nao encontrado")

    del perfis[perfil_id]
    _save_perfis(perfis)
    return {"status": "deleted", "id": perfil_id}
