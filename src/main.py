from datetime import datetime
import traceback

from src.pncp_api_client import fetch_licitacoes_from_pncp_api
from src.hermes_db import (
    init_db,
    start_run,
    finish_run,
    upsert_licitacoes,
    log_event,
    DB_PATH
)

# ============================================================
# CONFIG
# ============================================================
DEFAULT_CONFIG = {
    "keywords": ["limpeza", "higiene", "detergente", "papel"],
    "states": ["SP", "RJ", "MG"],
    "max_total": 50,
}

# ============================================================
# PIPELINE PRINCIPAL
# ============================================================
def run_pipeline(perfil: str, config: dict):
    print("============================================================")
    print("🚀 PROJETO HERMES PREMIUM")
    print("============================================================")

    init_db(DB_PATH)

    run_id = None

    try:
        # 🔥 INÍCIO
        run_id = start_run(db_path=DB_PATH)

        print("Coletando da API PNCP...")

        licitacoes, stats = fetch_licitacoes_from_pncp_api(
            keywords=config["keywords"],
            states=config["states"],
            max_total=config["max_total"],
            run_id=run_id,
            event_logger=log_event
        )

        # =========================
        # PASSO 2 - ENRIQUECIMENTO + INTELIGÊNCIA
        # =========================
        from src.hermes_score import calcular_score
        from src.hermes_db import get_price_reference

        def detectar_oportunidade(lic, referencia):
            valor = lic.get("valor_estimado_num", 0)
            media = referencia.get("media")

            if not media or media == 0:
                return {"oportunidade": "desconhecida", "delta": 0}

            delta = (valor - media) / media

            if delta > 0.5:
                nivel = "🔥 MUITO ACIMA (OURO)"
            elif delta > 0.2:
                nivel = "🟡 ACIMA"
            elif delta < -0.2:
                nivel = "🔵 ABAIXO"
            else:
                nivel = "⚖️ NORMAL"

            return {
                "oportunidade": nivel,
                "delta": round(delta, 2)
            }


        for lic in licitacoes:
            # 🔎 referência de preço histórico
            ref = get_price_reference(
                lic.get("keyword"),
                lic.get("estado")
            )

            # 🧠 score inteligente
            resultado = calcular_score(lic, ref)

            # 💰 oportunidade
            oportunidade = detectar_oportunidade(lic, ref)

            # 📦 enrich final
            lic["score"] = resultado.get("score", 0)
            lic["classificacao"] = resultado.get("classificacao", "")
            lic["motivos_score"] = resultado.get("motivos", [])

            lic["preco_medio_ref"] = ref.get("media")
            lic["oportunidade"] = oportunidade["oportunidade"]
            lic["delta_preco"] = oportunidade["delta"]

        # =========================
        # 🚨 ALERTA AUTOMÁTICO
        # =========================
        alertas = []

        for lic in licitacoes:
            score = lic.get("score", 0)
            delta = lic.get("delta_preco", 0)
            oportunidade = lic.get("oportunidade", "")

            # 🔥 REGRA DE OURO
            if score >= 5 and delta > 0.4:
                alerta = {
                    "pncp_id": lic.get("pncp_id"),
                    "objeto": lic.get("objeto"),
                    "valor": lic.get("valor_estimado_num"),
                    "estado": lic.get("estado"),
                    "oportunidade": oportunidade,
                    "score": score,
                    "delta": delta
                }

                alertas.append(alerta)

                print("\n🚨 OPORTUNIDADE FORTE DETECTADA!")
                print(f"Objeto: {alerta['objeto']}")
                print(f"Valor: {alerta['valor']}")
                print(f"Estado: {alerta['estado']}")
                print(f"Score: {score} | Delta: {delta}")
                print("--------------------------------------------------")

        print(f"Processando {len(licitacoes)} registros...")

        saved = upsert_licitacoes(licitacoes, db_path=DB_PATH)

        print(f"Salvos no banco: {saved}")

        # =========================
        # ENVIO DE E-MAIL AUTOMÁTICO
        # =========================
        from src.notifier import enviar_email_alerta

        enviar_email_alerta(alertas)    

        # =========================
        # PASSO 3 - EXPORTAÇÃO EXCEL (COM INTELIGÊNCIA)
        # =========================
        try:
            import pandas as pd
            from pathlib import Path
            from openpyxl.styles import PatternFill

            print("📦 Gerando Excel...")

            df = pd.DataFrame(licitacoes)

            if df.empty:
                print("⚠️ Nenhum dado para exportar.")
            else:
                # =========================
                # 🔥 RANKING INTELIGENTE
                # =========================
                df["ranking_final"] = (
                    df["score"].fillna(0) * 2 +
                    df["delta_preco"].fillna(0) * 5
                )

                # ordena pelo melhor
                df = df.sort_values(by="ranking_final", ascending=False)

                # =========================
                # 🎯 FILTRO DE OPORTUNIDADES
                # =========================
                df_oportunidades = df[
                    (df["score"] >= 4) &
                    (df["delta_preco"] > 0.2)
                ]

                # =========================
                # PRIORIDADE
                # =========================
                def classificar(score):
                    if score >= 5:
                        return "🔥 ALTA"
                    elif score >= 3:
                        return "⚠️ MÉDIA"
                    else:
                        return "BAIXA"

                df["prioridade"] = df["score"].apply(classificar)

                # =========================
                # SALVAR
                # =========================
                OUTPUT_DIR = Path("output")
                OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

                file_path = OUTPUT_DIR / f"hermes_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"

                with pd.ExcelWriter(file_path, engine="openpyxl") as writer:
                    # 📊 aba principal
                    df.to_excel(writer, index=False, sheet_name="Oportunidades")

                    # 🔥 aba inteligente
                    df_oportunidades.to_excel(writer, index=False, sheet_name="🔥 OPORTUNIDADES")

                    # 🏆 TOP 10
                    df.head(10).to_excel(writer, index=False, sheet_name="TOP 10")

                    # =========================
                    # 🎨 FORMATAÇÃO
                    # =========================
                    ws = writer.sheets["Oportunidades"]

                    verde = PatternFill(start_color="C6EFCE", fill_type="solid")
                    amarelo = PatternFill(start_color="FFF2CC", fill_type="solid")
                    vermelho = PatternFill(start_color="F8CBAD", fill_type="solid")

                    col_prioridade = None
                    for idx, col in enumerate(df.columns, start=1):
                        if col == "prioridade":
                            col_prioridade = idx
                            break

                    if col_prioridade:
                        for row in range(2, len(df) + 2):
                            cell = ws.cell(row=row, column=col_prioridade)
                            valor = str(cell.value)

                            if "ALTA" in valor:
                                cell.fill = verde
                            elif "MÉDIA" in valor:
                                cell.fill = amarelo
                            else:
                                cell.fill = vermelho

                    if alertas:
                        df_alertas = pd.DataFrame(alertas)
                        df_alertas.to_excel(writer, sheet_name="🚨 ALERTAS", index=False)

                print(f"✅ Excel salvo em: {file_path.resolve()}")

                print(f"🔥 Oportunidades reais encontradas: {len(df_oportunidades)}")

        except Exception as e:
            print(f"❌ ERRO AO GERAR EXCEL: {e}")


        # 🔥 FINALIZAÇÃO
        finish_run(
            run_id,
            stats,
            status="finished",
            db_path=DB_PATH
        )

        print("\n📊 RESUMO FINAL")
        print(f"Total coletado: {len(licitacoes)}")
        print(f"Stats: {stats}")

    except Exception as e:
        print("\n❌ ERRO FATAL NO HERMES\n")
        print(str(e))
        traceback.print_exc()

        if run_id:
            finish_run(
                run_id,
                {},
                status="error",
                notes=str(e),
                db_path=DB_PATH
            )


# ============================================================
# ENTRYPOINT
# ============================================================
def main():
    print("\n=== PROJETO HERMES ===")
    print("1. Rodar coleta PNCP")
    print("0. Sair")

    op = input("\nEscolha uma opção: ").strip()

    if op == "1":
        perfil = input("Perfil (ex: medicamentos): ").strip() or "default"
        config = DEFAULT_CONFIG
        run_pipeline(perfil, config)
    else:
        print("Encerrando...")


if __name__ == "__main__":
    main()